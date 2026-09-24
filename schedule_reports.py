"""Shared weekly Gantt and detailed reports; PDF is generated from current data."""
import datetime as dt
import html
import io

import pandas as pd

from commitment_schedule import BROWN, SOURCE, WEEK_NOTE, SUMMARY_NOTE, commitment_details, gantt_progress, source_rows, week_date, week_label, major_group

# ─── Màu sắc thương hiệu ──────────────────────────────────────────────────────
NAVY       = '#002C6C'
NAVY_LIGHT = '#1A4A8A'
SILVER     = '#E8EDF2'
BORDER     = '#B0BBC6'
TEXT_DARK  = '#172B3A'
TEXT_MUTED = '#586875'
GREEN_OK   = '#059669'
RED_LATE   = '#C83232'
AMBER      = '#D97706'

CSS = '''
@page {
    size: A4 landscape;
    margin: 10mm 9mm 13mm 9mm;
}
* { box-sizing: border-box; }
body {
    font-family: "Segoe UI", Arial, sans-serif;
    font-size: 10px;
    color: #172B3A;
    margin: 0; padding: 0;
    background: #fff;
}

/* ── HEADER ── */
.report-header {
    display: flex;
    align-items: flex-start;
    justify-content: space-between;
    border-bottom: 3px solid #002C6C;
    padding-bottom: 6px;
    margin-bottom: 8px;
}
.report-header .brand {
    display: flex;
    flex-direction: column;
}
.report-header h1 {
    font-size: 17px;
    font-weight: 800;
    color: #002C6C;
    margin: 0 0 2px;
    letter-spacing: .3px;
}
.report-header .sub-title {
    font-size: 9.5px;
    color: #586875;
    margin: 0;
}
.report-header .meta-box {
    text-align: right;
    font-size: 9px;
    color: #586875;
    line-height: 1.6;
}
.report-header .badge {
    display: inline-block;
    background: #002C6C;
    color: white;
    font-weight: bold;
    font-size: 8.5px;
    padding: 1px 6px;
    border-radius: 3px;
    margin-bottom: 2px;
}

/* ── SECTION TITLE ── */
h2 {
    font-size: 11px;
    font-weight: 700;
    color: #002C6C;
    margin: 6px 0 4px;
    text-transform: uppercase;
    letter-spacing: .4px;
}
p { margin: 2px 0 5px; }

/* ── LEGEND ── */
.legend {
    font-size: 8.5px;
    color: #586875;
    background: #F0F4F8;
    border-left: 3px solid #002C6C;
    padding: 3px 6px;
    margin-bottom: 6px;
    border-radius: 0 3px 3px 0;
}

/* ── GANTT TABLE ── */
.gantt {
    width: 100%;
    border-collapse: collapse;
    table-layout: fixed;
}
.gantt th {
    background: #002C6C;
    color: white;
    font-weight: 700;
    font-size: 8px;
    padding: 3px 2px;
    text-align: center;
    border: 1px solid #1A4A8A;
}
.gantt th.label-col {
    text-align: left;
    padding-left: 5px;
    font-size: 8.5px;
}
.gantt td {
    border: 1px solid #D1D9E0;
    padding: 0;
    height: 24px;
    vertical-align: middle;
}
.gantt td.label-col {
    padding: 4px 5px;
    font-size: 8.5px;
    vertical-align: top;
    line-height: 1.4;
    height: auto;
}
.gantt tr:nth-child(even) td.label-col { background: #F7F9FB; }
.gantt tr:nth-child(even) td           { background: #FAFBFC; }

/* ── DETAIL TABLE ── */
.detail {
    width: 100%;
    border-collapse: collapse;
    table-layout: fixed;
}
.detail th {
    background: #002C6C;
    color: white;
    font-weight: 700;
    font-size: 8.5px;
    padding: 4px 5px;
    border: 1px solid #1A4A8A;
    vertical-align: middle;
}
.detail td {
    font-size: 8.5px;
    line-height: 1.4;
    padding: 5px 5px;
    border: 1px solid #D1D9E0;
    vertical-align: top;
}
.detail tr:nth-child(even) td { background: #F4F7FA; }
.detail tr:hover td            { background: #EAF0F7; }

/* ── PROGRESS BAR ── */
.progress-wrap {
    background: #E2E8F0;
    border-radius: 3px;
    height: 7px;
    margin-top: 2px;
    overflow: hidden;
}
.progress-bar {
    height: 7px;
    border-radius: 3px;
}

/* ── STATUS BADGES ── */
.badge-done     { color: #059669; font-weight: 700; }
.badge-inprog   { color: #1A4A8A; font-weight: 700; }
.badge-late     { color: #C83232; font-weight: 700; }
.badge-pending  { color: #586875; }
.muted          { color: #586875; font-size: 8px; }

/* ── PAGE BREAK ── */
.page { page-break-after: always; margin-bottom: 22px; }
.page:last-child { page-break-after: auto; }
thead { display: table-header-group; }
tr { break-inside: avoid; }

/* ── FOOTER ── */
.footer {
    position: fixed;
    bottom: 6mm;
    left: 9mm; right: 9mm;
    font-size: 7.5px;
    color: #8A9BAB;
    display: flex;
    justify-content: space-between;
    border-top: 1px solid #D1D9E0;
    padding-top: 2px;
}

@media print {
    .no-print { display: none; }
    .page { margin-bottom: 0; }
    body { print-color-adjust: exact; -webkit-print-color-adjust: exact; }
    .footer { display: flex; }
}
@media screen {
    .footer { display: none; }
    body { padding: 12px; }
}
'''


def esc(value):
    return html.escape(str(value)).replace('\n', '<br>')


def is_summary(df):
    return not df.empty and df['Mã'].str.startswith('HM-').all()


def report_metadata(df):
    metadata = {r['code']: r for r in source_rows()}
    for _, r in df.iterrows():
        if str(r['Mã']).startswith('HM-'):
            weeks = [dict(year=d.year, month=d.month, week=min((d.day-1)//7+1,4)) for d in (r['Bắt đầu'],r['Hoàn thành'])]
            metadata[r['Mã']] = dict(current_weeks=weeks, group=major_group(r['Mã']))
    return metadata


def _status_badge(status_text):
    s = str(status_text).strip()
    if 'hoàn' in s.lower():
        return f'<span class="badge-done">✔ {esc(s)}</span>'
    if 'đang' in s.lower() or 'thi công' in s.lower():
        return f'<span class="badge-inprog">▶ {esc(s)}</span>'
    if 'quá' in s.lower() or 'hạn' in s.lower():
        return f'<span class="badge-late">⚠ {esc(s)}</span>'
    return f'<span class="badge-pending">{esc(s)}</span>'


def _progress_bar_html(pct, color=None):
    if color is None:
        color = GREEN_OK if pct >= 100 else (NAVY_LIGHT if pct >= 50 else AMBER)
    return (f'<b>{pct}%</b>'
            f'<div class="progress-wrap">'
            f'<div class="progress-bar" style="width:{min(pct,100)}%;background:{color}"></div>'
            f'</div>')


def _make_header_html(report_date, unit_count, unit_label, summary):
    badge_html = f'<span class="badge">PDF · A4 Ngang</span><br>'
    return f'''
<div class="report-header">
  <div class="brand">
    <h1>🏗 TIẾN ĐỘ DỰ ÁN — HYUNDAI MIỀN TÂY · CẦN THƠ</h1>
    <p class="sub-title">Chủ đầu tư: <b>Thế Giới Xe Tải</b> &nbsp;|&nbsp; Nguồn cam kết: <b>{esc(SOURCE)}</b>
      &nbsp;|&nbsp; Theo dõi: <b>{unit_count} {unit_label}</b>
      {"&nbsp;|&nbsp; <i>" + SUMMARY_NOTE + "</i>" if summary else ""}
    </p>
  </div>
  <div class="meta-box">
    {badge_html}
    Mốc theo dõi: <b>{report_date:%d/%m/%Y}</b><br>
    Xuất lúc: {dt.datetime.now():%H:%M %d/%m/%Y}
  </div>
</div>
'''


def report_sections(df, report_date):
    metadata = report_metadata(df)
    summary = is_summary(df)
    unit = 'hạng mục lớn' if summary else 'công việc'
    legend = (f'<div class="legend">'
              f'<b style="color:{BROWN}">■ Màu nâu:</b> Cam kết áp dụng từ mục 7. '
              f'Mục 15 tạm ẩn trên biểu đồ Gantt; vẫn có đầy đủ trong bảng chi tiết. '
              f'{WEEK_NOTE}'
              f'</div>')

    visible = gantt_progress(df)
    dated = visible.loc[[bool(metadata.get(c, {}).get('current_weeks', True)) for c in visible['Mã']]]
    first = min(pd.to_datetime(dated['Bắt đầu']).min().date(), dt.date(2026, 9, 1)).replace(day=1)
    last = max(pd.to_datetime(dated['Hoàn thành']).max().date(), dt.date(2026, 12, 31))
    months = []
    cursor = first
    while cursor <= last:
        months.append((cursor.year, cursor.month))
        cursor = dt.date(cursor.year+cursor.month//12, cursor.month%12+1, 1)
    weeks = [dict(year=y, month=m, week=w) for y, m in months for w in range(1, 5)]

    # Tên tháng tiếng Việt viết tắt
    vi_months = {1:'T1',2:'T2',3:'T3',4:'T4',5:'T5',6:'T6',7:'T7',8:'T8',9:'T9',10:'T10',11:'T11',12:'T12'}
    head = ('<thead>'
            '<tr>'
            '<th class="label-col" rowspan="2" style="width:33%">Mã / Hạng mục — Diễn giải</th>'
            + ''.join(f'<th colspan="4">{vi_months[m]}/{y}</th>' for y, m in months)
            + '</tr><tr>'
            + ''.join(f'<th>T{w["week"]}</th>' for w in weeks)
            + '</tr></thead>')

    sections = []
    per_page = 18
    for start in range(0, len(visible), per_page):
        chunk = visible.iloc[start:start+per_page]
        header_html = _make_header_html(report_date, len(df), unit, summary)
        lines = []
        for _, r in chunk.iterrows():
            code = r['Mã']
            item = metadata.get(code)
            label = (f'<b>{esc(code)}</b> · {esc(r["Hạng mục công việc"])}'
                     f'<br><span class="muted">📍 {esc(r["Phân khu"])}</span>')
            cells = []
            historical = item and not item['current_weeks']
            if historical:
                cells = [f'<td colspan="{len(weeks)}" style="color:{GREEN_OK};font-style:italic">'
                         f'✔ Đã hoàn thiện theo mốc cũ trong PDF cam kết; chưa có mốc mới.</td>']
            else:
                for w in weeks:
                    overlaps = r['Bắt đầu'] <= week_date(w, True) and r['Hoàn thành'] >= week_date(w)
                    color = BROWN if item and item.get('group', 7) >= 7 else GREEN_OK if r['Tiến độ (%)'] == 100 else '#64748B'
                    in_week = week_date(w) <= report_date <= week_date(w, True)
                    style = f'background:{color};' if overlaps else 'background:#F8FAFC;'
                    if in_week:
                        style += f'border-left:2px solid {RED_LATE};'
                    cells.append(f'<td style="{style}">&nbsp;</td>')
            lines.append(f'<tr><td class="label-col">{label}</td>{"".join(cells)}</tr>')

        pg_label = f'Công việc {start+1}–{min(start+per_page, len(visible))}/{len(visible)}'
        sections.append(
            header_html
            + f'<h2>📊 1. Biểu đồ Gantt theo tuần &nbsp;<small style="font-weight:400;color:{TEXT_MUTED}">({pg_label})</small></h2>'
            + legend
            + f'<table class="gantt">{head}<tbody>{"".join(lines)}</tbody></table>'
        )

    # Bảng chi tiết
    heading_key = '2. TỔNG HỢP HẠNG MỤC LỚN' if summary else '2. BẢNG CHI TIẾT'
    detail_head = (
        '<thead><tr>'
        '<th style="width:11%">Mã / Phân khu</th>'
        '<th style="width:22%">Diễn giải công việc</th>'
        '<th style="width:17%">Mốc cam kết theo tuần</th>'
        '<th style="width:11%">Thực tế / Tiến độ</th>'
        '<th style="width:8%">P.I.C</th>'
        '<th style="width:31%">Nguồn & Ghi chú</th>'
        '</tr></thead>'
    )
    for start in range(0, len(df), 10):
        chunk = df.iloc[start:start+10]
        header_html = _make_header_html(report_date, len(df), unit, summary)
        lines = []
        for _, r in chunk.iterrows():
            item = metadata.get(r['Mã'])
            current = item['current_weeks'] if item else []
            if current:
                period = f'{week_label(current[0])}<br>→ {week_label(current[-1])}'
            elif item:
                period = '<i style="color:#D97706">Chỉ có mốc cũ trên PDF</i>'
            else:
                period = f'<span class="muted">Ngoài phạm vi mục 7–15</span>'
            period += (f'<br><span class="muted">📅 {r["Bắt đầu"]:%d/%m/%Y}'
                       f' – {r["Hoàn thành"]:%d/%m/%Y}</span>')

            pct = int(r['Tiến độ (%)'])
            bar_color = GREEN_OK if pct >= 100 else (NAVY_LIGHT if pct >= 50 else AMBER)
            progress_html = _progress_bar_html(pct, bar_color)
            status_html = _status_badge(r['Trạng thái'])

            values = [
                f'<b>{esc(r["Mã"])}</b><br><span class="muted">📍 {esc(r["Phân khu"])}</span>',
                esc(r['Hạng mục công việc']),
                period,
                f'{progress_html}<br>{status_html}',
                esc(r['Người phụ trách']),
                esc(r['Ghi chú']),
            ]
            lines.append('<tr>' + ''.join(f'<td>{v}</td>' for v in values) + '</tr>')

        pg_label = f'{start+1}–{min(start+10, len(df))}/{len(df)}'
        sections.append(
            header_html
            + f'<h2>📋 {heading_key} &nbsp;<small style="font-weight:400;color:{TEXT_MUTED}">({pg_label})</small></h2>'
            + f'<p class="muted">⚡ {WEEK_NOTE} &nbsp; Tiến độ thực tế không suy ra từ độ dài thanh kế hoạch.</p>'
            + f'<table class="detail">{detail_head}<tbody>{"".join(lines)}</tbody></table>'
        )

    return sections


def build_printable_html(df, report_date):
    sections = report_sections(df, report_date)
    footer = (f'<div class="footer">'
              f'<span>🏗 Dự án Hyundai Miền Tây – Cần Thơ &nbsp;|&nbsp; Chủ đầu tư: Thế Giới Xe Tải</span>'
              f'<span>Xuất: {dt.datetime.now():%H:%M %d/%m/%Y} &nbsp;|&nbsp; Nguồn: {esc(SOURCE)}</span>'
              f'</div>')
    body_sections = ''.join(f'<section class="page">{s}</section>' for s in sections)
    return (f'<!doctype html><html lang="vi"><head><meta charset="utf-8">'
            f'<title>Tiến độ Hyundai Miền Tây – Cần Thơ {report_date:%d/%m/%Y}</title>'
            f'<style>{CSS}</style></head><body>'
            f'<button class="no-print" onclick="window.print()" '
            f'style="margin-bottom:10px;padding:8px 20px;background:#002C6C;color:white;border:none;'
            f'border-radius:5px;font-size:13px;cursor:pointer">🖨 In báo cáo</button>'
            f'{body_sections}'
            f'{footer}'
            f'</body></html>')


# ─── PDF bằng PyMuPDF ─────────────────────────────────────────────────────────

def build_pdf(df, report_date):
    import fitz
    doc = fitz.open()
    metadata = report_metadata(df)
    summary = is_summary(df)

    def color(h):
        h = h.lstrip('#')
        return tuple(int(h[i:i+2], 16)/255 for i in (0, 2, 4))

    NAVY_RGB   = color(NAVY)
    SILVER_RGB = color(SILVER)
    BORDER_RGB = color(BORDER)
    WHITE_RGB  = (1.0, 1.0, 1.0)
    BROWN_RGB  = color(BROWN)
    GREEN_RGB  = color(GREEN_OK)
    AMBER_RGB  = color(AMBER)
    SLATE_RGB  = color('#64748B')
    RED_RGB    = color(RED_LATE)
    ZEBRA_RGB  = color('#F4F7FA')

    # ── Layout constants (A4 landscape: 842×595 pt) ───────────────────────────
    PW, PH = 842, 595
    MARGIN_L, MARGIN_R = 24, 24
    MARGIN_T, MARGIN_B = 18, 30        # bottom reserved for footer
    CONTENT_W = PW - MARGIN_L - MARGIN_R

    def draw_rect_filled(page, rect, fill_rgb, stroke_rgb=None, stroke_w=0.4):
        page.draw_rect(fitz.Rect(rect), color=stroke_rgb or BORDER_RGB,
                       fill=fill_rgb, width=stroke_w)

    def insert_html(page, rect, content, font_size=8):
        wrapped = (f'<div style="font-family:\'Segoe UI\',Arial,sans-serif;'
                   f'font-size:{font_size}px;color:#172B3A;line-height:1.4">{content}</div>')
        page.insert_htmlbox(fitz.Rect(rect) + (3, 2, -3, -2), wrapped)

    def draw_header(page):
        """Vẽ header chuẩn cho mỗi trang."""
        # Đường kẻ ngang navy trên cùng
        page.draw_rect(fitz.Rect(MARGIN_L, MARGIN_T, PW - MARGIN_R, MARGIN_T + 3),
                       color=NAVY_RGB, fill=NAVY_RGB, width=0)
        # Tiêu đề
        title_html = (
            f'<b style="font-size:15px;color:{NAVY}">🏗 TIẾN ĐỘ DỰ ÁN — HYUNDAI MIỀN TÂY · CẦN THƠ</b><br>'
            f'<span style="font-size:8px;color:{TEXT_MUTED}">'
            f'Chủ đầu tư: <b>Thế Giới Xe Tải</b> &nbsp;·&nbsp; '
            f'Nguồn cam kết: <b>{esc(SOURCE)}</b>'
            + (f' &nbsp;·&nbsp; <i>{SUMMARY_NOTE}</i>' if summary else '')
            + f'</span>'
        )
        page.insert_htmlbox(fitz.Rect(MARGIN_L, MARGIN_T + 4, CONTENT_W * 0.72, MARGIN_T + 38), title_html)
        # Meta box phía phải
        meta_html = (
            f'<div style="text-align:right;font-size:8px;color:{TEXT_MUTED};line-height:1.6">'
            f'<b style="background:{NAVY};color:white;padding:1px 5px;border-radius:2px">PDF · A4 Ngang</b><br>'
            f'Mốc theo dõi: <b>{report_date:%d/%m/%Y}</b><br>'
            f'Xuất: {dt.datetime.now():%H:%M %d/%m/%Y}'
            f'</div>'
        )
        page.insert_htmlbox(fitz.Rect(CONTENT_W * 0.72 + MARGIN_L, MARGIN_T + 4, PW - MARGIN_R, MARGIN_T + 38), meta_html)
        # Đường kẻ phân cách header/nội dung
        page.draw_line((MARGIN_L, MARGIN_T + 40), (PW - MARGIN_R, MARGIN_T + 40),
                       color=NAVY_RGB, width=1)
        return MARGIN_T + 44   # y bắt đầu nội dung

    def draw_footer(page, page_num, total_pages):
        y = PH - MARGIN_B + 4
        try:
            page.draw_line((MARGIN_L, y), (PW - MARGIN_R, y), color=BORDER_RGB, width=0.5)
        except Exception:
            pass
        left_txt = f'Dự án Hyundai Miền Tây – Cần Thơ  ·  Chủ đầu tư: Thế Giới Xe Tải'
        right_txt = f'Trang {page_num}/{total_pages}  ·  Nguồn: {SOURCE}'
        page.insert_text((MARGIN_L, y + 9), left_txt,
                         fontsize=7, color=color('#8A9BAB'))
        page.insert_text((PW - MARGIN_R - 160, y + 9), right_txt,
                         fontsize=7, color=color('#8A9BAB'))

    # ── Gantt ─────────────────────────────────────────────────────────────────
    visible = gantt_progress(df)
    dated = visible.loc[[bool(metadata.get(c, {}).get('current_weeks', True)) for c in visible['Mã']]]
    first = min(pd.to_datetime(dated['Bắt đầu']).min().date(), dt.date(2026, 9, 1)).replace(day=1)
    last  = max(pd.to_datetime(dated['Hoàn thành']).max().date(), dt.date(2026, 12, 31))
    months = []
    cursor = first
    while cursor <= last:
        months.append((cursor.year, cursor.month))
        cursor = dt.date(cursor.year + cursor.month // 12, cursor.month % 12 + 1, 1)
    weeks = [dict(year=y, month=m, week=w) for y, m in months for w in range(1, 5)]

    LABEL_W   = 290
    cell_w    = (CONTENT_W - LABEL_W) / len(weeks)
    HEADER_H1 = 16   # hàng tháng
    HEADER_H2 = 14   # hàng tuần
    ROW_H     = 25   # mỗi hàng công việc

    PER_PAGE_GANTT = 16
    PER_PAGE_DET = 10
    total_gantt_pages = max(1, (len(visible) + PER_PAGE_GANTT - 1) // PER_PAGE_GANTT)
    total_detail_pages = max(1, (len(df) + PER_PAGE_DET - 1) // PER_PAGE_DET)
    total_pages = total_gantt_pages + total_detail_pages

    for page_idx_gantt, start in enumerate(range(0, len(visible), PER_PAGE_GANTT)):
        page = doc.new_page(width=PW, height=PH)
        y0 = draw_header(page)
        draw_footer(page, page_idx_gantt + 1, total_pages)

        # Sub-title
        chunk_label = f'1. BIỂU ĐỒ GANTT · Công việc {start+1}–{min(start+PER_PAGE_GANTT,len(visible))}/{len(visible)}'
        page.insert_htmlbox(
            fitz.Rect(MARGIN_L, y0, PW - MARGIN_R, y0 + 12),
            f'<b style="font-size:10px;color:{NAVY}">{chunk_label}</b>'
            f'<span style="font-size:7.5px;color:{TEXT_MUTED}"> &nbsp;·&nbsp; '
            f'■ Nâu = cam kết áp dụng. Mục 15 tạm ẩn trên biểu đồ.</span>'
        )
        y0 += 14

        x0 = MARGIN_L
        # Header cột nhãn
        draw_rect_filled(page, (x0, y0, x0 + LABEL_W, y0 + HEADER_H1 + HEADER_H2), NAVY_RGB)
        page.insert_htmlbox(
            fitz.Rect(x0 + 3, y0 + 3, x0 + LABEL_W - 3, y0 + HEADER_H1 + HEADER_H2 - 3),
            f'<b style="font-size:8.5px;color:white">Mã / Hạng mục — Diễn giải &amp; Phân khu</b>'
        )
        # Header tháng + tuần
        for i, (y, m) in enumerate(months):
            x = x0 + LABEL_W + i * 4 * cell_w
            draw_rect_filled(page, (x, y0, x + 4 * cell_w, y0 + HEADER_H1), NAVY_RGB, NAVY_RGB)
            page.insert_htmlbox(
                fitz.Rect(x + 1, y0 + 1, x + 4 * cell_w - 1, y0 + HEADER_H1 - 1),
                f'<b style="font-size:8px;color:white;text-align:center">{m:02}/{y}</b>'
            )
        for i, w in enumerate(weeks):
            x = x0 + LABEL_W + i * cell_w
            draw_rect_filled(page, (x, y0 + HEADER_H1, x + cell_w, y0 + HEADER_H1 + HEADER_H2),
                             color(NAVY_LIGHT), color(NAVY_LIGHT))
            page.insert_htmlbox(
                fitz.Rect(x + 1, y0 + HEADER_H1 + 1, x + cell_w - 1, y0 + HEADER_H1 + HEADER_H2 - 1),
                f'<span style="font-size:7px;color:white">T{w["week"]}</span>'
            )

        y0 += HEADER_H1 + HEADER_H2

        for idx, (_, r) in enumerate(visible.iloc[start:start + PER_PAGE_GANTT].iterrows()):
            ry = y0 + idx * ROW_H
            zebra = ZEBRA_RGB if idx % 2 == 0 else WHITE_RGB
            code = r['Mã']
            item = metadata.get(code)

            # Label cell
            draw_rect_filled(page, (x0, ry, x0 + LABEL_W, ry + ROW_H), zebra)
            label_html = (f'<b style="font-size:8px">{esc(code)}</b>'
                          f' <span style="font-size:8px">· {esc(r["Hạng mục công việc"])}</span>'
                          f'<br><span style="font-size:7px;color:{TEXT_MUTED}">📍 {esc(r["Phân khu"])}</span>')
            page.insert_htmlbox(fitz.Rect(x0 + 3, ry + 2, x0 + LABEL_W - 3, ry + ROW_H - 2), label_html)

            # Bar cells
            historical = item and not item['current_weeks']
            if historical:
                draw_rect_filled(page, (x0 + LABEL_W, ry, PW - MARGIN_R, ry + ROW_H), zebra)
                page.insert_htmlbox(
                    fitz.Rect(x0 + LABEL_W + 4, ry + 4, PW - MARGIN_R - 4, ry + ROW_H - 4),
                    f'<span style="color:{GREEN_OK};font-size:7.5px;font-style:italic">'
                    f'✔ Đã hoàn thiện theo mốc cũ; chưa có mốc mới.</span>'
                )
            else:
                fill_color = (BROWN_RGB if item and item.get('group', 7) >= 7
                              else GREEN_RGB if r['Tiến độ (%)'] == 100
                              else SLATE_RGB)
                for j, w in enumerate(weeks):
                    wx = x0 + LABEL_W + j * cell_w
                    overlaps = r['Bắt đầu'] <= week_date(w, True) and r['Hoàn thành'] >= week_date(w)
                    cell_bg = fill_color if overlaps else zebra
                    draw_rect_filled(page, (wx, ry, wx + cell_w, ry + ROW_H), cell_bg, BORDER_RGB, 0.3)
                    if week_date(w) <= report_date <= week_date(w, True):
                        page.draw_line((wx, ry), (wx, ry + ROW_H), color=RED_RGB, width=1.2)

    # ── Bảng chi tiết ─────────────────────────────────────────────────────────
    COL_WIDTHS = [82, 168, 135, 90, 65, 254]
    COL_LABELS = ['Mã / Phân khu', 'Diễn giải công việc',
                  'Mốc cam kết theo tuần', 'Thực tế / Tiến độ', 'P.I.C', 'Nguồn & Ghi chú']
    HEADER_H   = 22
    ROW_H_DET  = 46

    PER_PAGE_DET = 10
    heading_key = '2. TỔNG HỢP HẠNG MỤC LỚN' if summary else '2. BẢNG CHI TIẾT'
    for page_idx_det, start in enumerate(range(0, len(df), PER_PAGE_DET)):
        page = doc.new_page(width=PW, height=PH)
        y0 = draw_header(page)
        draw_footer(page, total_gantt_pages + page_idx_det + 1, total_pages)

        chunk_label = f'{heading_key} · {start+1}–{min(start+PER_PAGE_DET,len(df))}/{len(df)}'
        page.insert_htmlbox(
            fitz.Rect(MARGIN_L, y0, PW - MARGIN_R, y0 + 12),
            f'<b style="font-size:10px;color:{NAVY}">{chunk_label}</b>'
            f'<span style="font-size:7.5px;color:{TEXT_MUTED}"> &nbsp;·&nbsp; '
            f'Bao gồm đầy đủ mục 15. Tiến độ thực tế không suy ra từ thanh Gantt.</span>'
        )
        y0 += 14

        # Header hàng
        x = MARGIN_L
        for label, w in zip(COL_LABELS, COL_WIDTHS):
            draw_rect_filled(page, (x, y0, x + w, y0 + HEADER_H), NAVY_RGB)
            page.insert_htmlbox(
                fitz.Rect(x + 3, y0 + 3, x + w - 3, y0 + HEADER_H - 3),
                f'<b style="font-size:8px;color:white">{label}</b>'
            )
            x += w
        y0 += HEADER_H

        for idx, (_, r) in enumerate(df.iloc[start:start + PER_PAGE_DET].iterrows()):
            ry = y0 + idx * ROW_H_DET
            zebra = ZEBRA_RGB if idx % 2 == 0 else WHITE_RGB
            item = metadata.get(r['Mã'])
            current = item['current_weeks'] if item else []
            if current:
                period = f'{week_label(current[0])} → {week_label(current[-1])}'
            elif item:
                period = '<i>Chỉ có mốc cũ trên PDF</i>'
            else:
                period = 'Ngoài phạm vi mục 7–15'
            period += f'<br><span style="font-size:7px;color:{TEXT_MUTED}">📅 {r["Bắt đầu"]:%d/%m/%Y} – {r["Hoàn thành"]:%d/%m/%Y}</span>'

            pct = int(r['Tiến độ (%)'])
            bar_color_rgb = GREEN_RGB if pct >= 100 else (color(NAVY_LIGHT) if pct >= 50 else AMBER_RGB)
            progress_html = (f'<b style="font-size:9px">{pct}%</b><br>'
                             f'<div style="background:#E2E8F0;border-radius:3px;height:6px;margin:2px 0;overflow:hidden">'
                             f'<div style="width:{min(pct,100)}%;height:6px;background:{"#"+NAVY_LIGHT[1:]};border-radius:3px"></div>'
                             f'</div>')

            status = str(r['Trạng thái'])
            if 'hoàn' in status.lower():
                status_html = f'<span style="color:{GREEN_OK};font-weight:bold;font-size:7.5px">✔ {esc(status)}</span>'
            elif 'đang' in status.lower():
                status_html = f'<span style="color:{NAVY_LIGHT};font-weight:bold;font-size:7.5px">▶ {esc(status)}</span>'
            else:
                status_html = f'<span style="color:{TEXT_MUTED};font-size:7.5px">{esc(status)}</span>'

            values = [
                f'<b style="font-size:8px">{esc(r["Mã"])}</b><br><span style="font-size:7px;color:{TEXT_MUTED}">📍 {esc(r["Phân khu"])}</span>',
                f'<span style="font-size:8px">{esc(r["Hạng mục công việc"])}</span>',
                f'<span style="font-size:7.5px">{period}</span>',
                progress_html + status_html,
                f'<span style="font-size:8px">{esc(r["Người phụ trách"])}</span>',
                f'<span style="font-size:7.5px">{esc(r["Ghi chú"])}</span>',
            ]
            x = MARGIN_L
            for value, w in zip(values, COL_WIDTHS):
                draw_rect_filled(page, (x, ry, x + w, ry + ROW_H_DET), zebra)
                page.insert_htmlbox(fitz.Rect(x + 3, ry + 3, x + w - 3, ry + ROW_H_DET - 3), value)
                x += w

    doc.subset_fonts()
    content = doc.tobytes(garbage=4, deflate=True)
    doc.close()
    return content


def build_excel(df, qcvn):
    from openpyxl.styles import PatternFill, Font, Alignment
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        (df if is_summary(df) else commitment_details(df)).to_excel(
            writer,
            sheet_name='Tong_Hop_Hang_Muc' if is_summary(df) else 'Tien_Do_Thi_Cong',
            index=False
        )
        qcvn.to_excel(writer, sheet_name='QCVN_121_Checklist', index=False)
        for sheet in writer.book:
            sheet.freeze_panes = 'C2'
            sheet.auto_filter.ref = sheet.dimensions
            for c in sheet[1]:
                c.fill = PatternFill('solid', fgColor='002C6C')
                c.font = Font(color='FFFFFF', bold=True)
            for row in sheet.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(vertical='top', wrap_text=True)
                if sheet.title == 'Tien_Do_Thi_Cong' and str(row[0].value).startswith('CK-'):
                    row[0].font = Font(color=BROWN[1:], bold=True)
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width = 24
            sheet.column_dimensions['B'].width = 48
            sheet.column_dimensions['I'].width = 90
    return output.getvalue()
